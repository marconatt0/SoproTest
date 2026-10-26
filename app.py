import datetime as dt
import difflib
import hashlib
import io
import re
import unicodedata
import xml.etree.ElementTree as ET
from html.parser import HTMLParser

import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import plotly.io as pio
import openpyxl

st.set_page_config(page_title="Gestão de Chamados SOPRO", page_icon="📊", layout="wide")

# Tema de todos os gráficos: legenda embaixo (não rouba largura no celular), arrastar o dedo rola a
# página em vez de dar zoom sem querer (tocar continua mostrando os valores) e margens enxutas.
pio.templates['sopro'] = go.layout.Template(layout=dict(
    dragmode=False, font=dict(size=13), hoverlabel=dict(font_size=14),
    legend=dict(orientation='h', yref='container', yanchor='bottom', y=0, xanchor='left', x=0),
    margin=dict(l=8, r=8, t=60, b=8), xaxis=dict(automargin=True), yaxis=dict(automargin=True)))
pio.templates.default = 'plotly+sopro'


def encurtar_nomes(fig, nomes, eixo='y'):
    """Nomes longos (empresas) com até 20 letras no eixo; o nome completo continua no valor ao tocar/passar o mouse."""
    curtos = [n if len(n) <= 20 else n[:19] + '…' for n in nomes]
    (fig.update_yaxes if eixo == 'y' else fig.update_xaxes)(tickvals=list(nomes), ticktext=curtos)


def altura_barras(n):
    """Altura de um gráfico de barras deitadas com n categorias."""
    return max(320, 28 * n + 140)

# ---------------------------------------------------------------------------
# Classificação dos status do portal
# ---------------------------------------------------------------------------
# Status do portal (comparados inteiros, sem acento/maiúscula) e a situação de cada um
STATUS_CATEGORIA = {"em andamento": "Aberto", "concluido": "Resolvido", "cancelado pela sharing": "Cancelado"}
# Status que não entram em nenhum cálculo nem relatório
STATUS_EXCLUIDOS = ["aguardando número do chamado", "cancelado pela claro", "fechado pelo mso", "pendente claro"]

CATEGORIA_CORES = {"Resolvido": "#2ECC71", "Aberto": "#FFA500", "Cancelado": "#95A5A6"}
SLA_CORES = {"NO PRAZO": "#2ECC71", "FORA DO PRAZO": "#E74C3C", "SEM SLA": "#BDC3C7"}


def normalizar_texto(texto):
    """Minúsculas, sem acentos e com qualquer pontuação/espaço (inclusive invisível) virando um espaço só."""
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(c for c in texto if not unicodedata.combining(c)).lower()
    return re.sub(r'[^a-z0-9]+', ' ', texto).strip()


STATUS_EXCLUIDOS_NORM = [normalizar_texto(s) for s in STATUS_EXCLUIDOS]


def status_excluido(status):
    s = normalizar_texto(status)
    return any(k in s for k in STATUS_EXCLUIDOS_NORM)


def categorizar_status(status):
    """Situação pela lista fechada; status desconhecido conta como Aberto (e é avisado na tela)."""
    return STATUS_CATEGORIA.get(normalizar_texto(status), "Aberto")


def status_desconhecidos(df):
    return sorted({s for s in df['Status'].unique() if normalizar_texto(s) not in STATUS_CATEGORIA})


CORES_NOMEADAS = {'red': 'FF0000', 'green': '008000', 'lime': '00FF00', 'darkgreen': '006400',
                  'darkred': '8B0000', 'crimson': 'DC143C', 'forestgreen': '228B22'}


def classificar_cor(cor_hex):
    """Converte a cor da bolinha de SLA (hex, ex.: 'FFFF0000' ou '#00B050') em NO PRAZO / FORA DO PRAZO."""
    if not cor_hex:
        return "SEM SLA"
    cor_hex = CORES_NOMEADAS.get(str(cor_hex).strip().lower(), str(cor_hex)).upper().lstrip('#')
    # Se a cor em HEX (RGB) tiver código de Verde (usado pelo Excel)
    if "00B050" in cor_hex or "00FF00" in cor_hex or "2ECC71" in cor_hex or "008000" in cor_hex:
        return "NO PRAZO"
    # Se a cor tiver código de Vermelho
    if "FF0000" in cor_hex or "C00000" in cor_hex or "FF4500" in cor_hex or "E74C3C" in cor_hex:
        return "FORA DO PRAZO"
    # Outros tons: decide pelo canal dominante
    if re.fullmatch(r'[0-9A-F]{6}|[0-9A-F]{8}', cor_hex):
        r, g, b = (int(cor_hex[-6:][i:i + 2], 16) for i in (0, 2, 4))
        if g >= 100 and g > r * 1.3 and g > b * 1.3:
            return "NO PRAZO"
        if r >= 150 and r > g * 1.8 and r > b * 1.8:
            return "FORA DO PRAZO"
    return "SEM SLA"


def cor_celula(cell):
    """Cor (hex) da fonte da célula do .xlsx ou, se não houver, do preenchimento."""
    if cell.font and cell.font.color and cell.font.color.type == 'rgb':
        return str(cell.font.color.rgb).upper()
    if cell.fill and cell.fill.fgColor and cell.fill.fgColor.type == 'rgb':
        return str(cell.fill.fgColor.rgb).upper()
    return None


# ---------------------------------------------------------------------------
# Leitura da planilha: .xlsx e os três formatos que podem vir com extensão .xls
# ---------------------------------------------------------------------------
def _ler_xlsx(conteudo):
    """.xlsx lido numa grade (valor + cor de cada célula), como os .xls: o cabeçalho é achado pela coluna
    "Estação" em qualquer linha e datas em texto ("29/set/26") são convertidas."""
    ws = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True).active
    valores, cores = [], []
    for linha in ws.iter_rows():
        valores.append([c.value for c in linha])
        cores.append([cor_celula(c) if c.value is not None else None for c in linha])
    return _grade_para_dataframe(valores, cores)


MESES_ABREVIADOS = {'jan': '01', 'fev': '02', 'mar': '03', 'abr': '04', 'mai': '05', 'jun': '06',
                    'jul': '07', 'ago': '08', 'set': '09', 'out': '10', 'nov': '11', 'dez': '12'}
MESES_ABREVIADOS_RE = re.compile(r'/(' + '|'.join(MESES_ABREVIADOS) + r')[a-zç]*\.?/', re.IGNORECASE)


def _grade_para_dataframe(valores, cores):
    """Monta o DataFrame a partir de uma grade (lista de linhas) e da cor de cada célula.

    O cabeçalho é a linha que tem a coluna "Estação"; o SLA vem da cor da coluna "SLA".
    """
    linha_cab = next((i for i, linha in enumerate(valores[:20])
                      if any(normalizar_texto(v) == 'estacao' for v in linha if v is not None)), None)
    if linha_cab is None:
        raise ValueError("Não encontrei a linha de cabeçalho (coluna 'Estação') na planilha.")

    nomes, vistos = [], {}
    for i, v in enumerate(valores[linha_cab]):
        nome = str(v).strip() if v not in (None, '') else f'Coluna {i + 1}'
        if nome in vistos:  # mesmo padrão do pandas: a 2ª coluna "Data" vira "Data.1"
            vistos[nome] += 1
            nome = f'{nome}.{vistos[nome]}'
        else:
            vistos[nome] = 0
        nomes.append(nome)

    col_sla = next((i for i, n in enumerate(nomes) if normalizar_texto(n) == 'sla'), None)
    linhas = valores[linha_cab + 1:]
    largura = len(nomes)
    linhas = [(list(l) + [None] * largura)[:largura] for l in linhas]
    df = pd.DataFrame(linhas, columns=nomes).replace({'': None})

    sla = []
    for i in range(linha_cab + 1, len(valores)):
        linha_cores = cores[i] if i < len(cores) else []
        cor = linha_cores[col_sla] if col_sla is not None and col_sla < len(linha_cores) else None
        sla.append(classificar_cor(cor))

    # Datas que chegam como texto (HTML/XML) ficam no formato dia/mês/ano.
    # O portal também usa o mês abreviado em português e ano com 2 dígitos (ex.: "13/out/26").
    for col in df.columns:
        textos = df[col].map(lambda v: isinstance(v, str))
        if normalizar_texto(col).startswith('data') and textos.any():
            datas = df.loc[textos, col].str.strip().str.replace(
                MESES_ABREVIADOS_RE, lambda m: f"/{MESES_ABREVIADOS[m.group(1).lower()]}/", regex=True)
            convertidas = pd.to_datetime(datas, dayfirst=True, errors='coerce', format='mixed')
            df[col] = df[col].where(~textos, convertidas)
    return df, sla


def _ler_xls_binario(conteudo):
    """.xls do Excel 97-2003 (formato binário). A cor vem da fonte ou do preenchimento da célula."""
    import xlrd
    livro = xlrd.open_workbook(file_contents=conteudo, formatting_info=True)
    folha = livro.sheet_by_index(0)

    def cor_rgb(indice):
        rgb = livro.colour_map.get(indice)
        return '%02X%02X%02X' % rgb if rgb else None

    valores, cores = [], []
    for r in range(folha.nrows):
        linha_v, linha_c = [], []
        for c in range(folha.ncols):
            celula = folha.cell(r, c)
            if celula.ctype == xlrd.XL_CELL_DATE:
                valor = xlrd.xldate_as_datetime(celula.value, livro.datemode)
            elif celula.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                valor = None
            else:
                valor = celula.value
            xf = livro.xf_list[folha.cell_xf_index(r, c)]
            cor = cor_rgb(livro.font_list[xf.font_index].colour_index)
            if classificar_cor(cor) == "SEM SLA":
                cor = cor_rgb(xf.background.pattern_colour_index) or cor
            linha_v.append(valor)
            linha_c.append(cor)
        valores.append(linha_v)
        cores.append(linha_c)
    return _grade_para_dataframe(valores, cores)


class _LeitorTabelaHTML(HTMLParser):
    """Lê a primeira tabela de um .xls que na verdade é uma página HTML (comum em exportações de portais)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.estilos_classe = {}
        self.linhas, self.cores = [], []
        self._linha = self._cores_linha = None
        self._celula = None
        self._em_style = False
        self._css = ''
        self._tabelas = 0
        self._cor_fonte = []

    @staticmethod
    def _cor_do_estilo(estilo):
        m = re.search(r'(?<![-\w])color\s*:\s*([#\w]+)', estilo or '', re.IGNORECASE)
        return m.group(1) if m else None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == 'style':
            self._em_style = True
        elif tag == 'table':
            self._tabelas += 1
        elif self._tabelas != 1:
            return
        elif tag == 'tr':
            self._linha, self._cores_linha = [], []
        elif tag in ('td', 'th') and self._linha is not None:
            cor = self._cor_do_estilo(a.get('style'))
            for classe in (a.get('class') or '').split():
                cor = cor or self.estilos_classe.get(classe)
            self._celula = {'texto': [], 'cor': cor, 'span': int(a.get('colspan') or 1)}
        elif tag == 'font' and self._celula is not None:
            self._celula['cor'] = a.get('color') or self._cor_do_estilo(a.get('style')) or self._celula['cor']
        elif tag in ('span', 'b', 'div') and self._celula is not None:
            self._celula['cor'] = self._cor_do_estilo(a.get('style')) or self._celula['cor']
        elif tag == 'br' and self._celula is not None:
            self._celula['texto'].append(' ')

    def handle_endtag(self, tag):
        if tag == 'style':
            self._em_style = False
            for seletor, corpo in re.findall(r'\.([\w-]+)\s*\{([^}]*)\}', self._css):
                cor = self._cor_do_estilo(corpo)
                if cor:
                    self.estilos_classe[seletor] = cor
        elif self._tabelas != 1:
            return
        elif tag in ('td', 'th') and self._celula is not None:
            texto = re.sub(r'\s+', ' ', ''.join(self._celula['texto'])).strip()
            for k in range(self._celula['span']):
                self._linha.append(texto if k == 0 else None)
                self._cores_linha.append(self._celula['cor'])
            self._celula = None
        elif tag == 'tr' and self._linha is not None:
            self.linhas.append(self._linha)
            self.cores.append(self._cores_linha)
            self._linha = None
        elif tag == 'table':
            self._tabelas += 1  # só a primeira tabela interessa

    def handle_data(self, data):
        if self._em_style:
            self._css += data
        elif self._celula is not None:
            self._celula['texto'].append(data)


def _ler_xls_html(texto):
    leitor = _LeitorTabelaHTML()
    leitor.feed(texto)
    valores = [[(v if v != '' else None) for v in linha] for linha in leitor.linhas]
    return _grade_para_dataframe(valores, leitor.cores)


def _ler_xls_xml(texto):
    """.xls no formato "Planilha XML 2003" do Excel (SpreadsheetML)."""
    ns = {'ss': 'urn:schemas-microsoft-com:office:spreadsheet'}
    q = lambda nome: '{urn:schemas-microsoft-com:office:spreadsheet}' + nome
    raiz = ET.fromstring(texto.encode('utf-8') if isinstance(texto, str) else texto)

    cor_estilo = {}
    for estilo in raiz.iter(q('Style')):
        fonte, interior = estilo.find('ss:Font', ns), estilo.find('ss:Interior', ns)
        cor = fonte.get(q('Color')) if fonte is not None else None
        if classificar_cor(cor) == "SEM SLA" and interior is not None:
            cor = interior.get(q('Color')) or cor
        cor_estilo[estilo.get(q('ID'))] = cor

    tabela = next(raiz.iter(q('Table')))
    valores, cores = [], []
    for linha in tabela.iter(q('Row')):
        indice_linha = linha.get(q('Index'))
        while indice_linha and len(valores) < int(indice_linha) - 1:
            valores.append([])
            cores.append([])
        lv, lc = [], []
        for celula in linha.findall('ss:Cell', ns):
            indice = celula.get(q('Index'))
            while indice and len(lv) < int(indice) - 1:
                lv.append(None)
                lc.append(None)
            dado = celula.find('ss:Data', ns)
            valor = ''.join(dado.itertext()).strip() if dado is not None else None
            if dado is not None and dado.get(q('Type')) == 'Number':
                valor = float(valor)
            elif dado is not None and dado.get(q('Type')) == 'DateTime':
                try:
                    valor = dt.datetime.fromisoformat(valor)
                except ValueError:
                    valor = pd.to_datetime(valor, errors='coerce')
            extra = int(celula.get(q('MergeAcross')) or 0)
            lv += [valor or None] + [None] * extra
            lc += [cor_estilo.get(celula.get(q('StyleID')))] * (1 + extra)
        valores.append(lv)
        cores.append(lc)
    return _grade_para_dataframe(valores, cores)


def ler_planilha(conteudo):
    """Identifica o formato pelo conteúdo (não pela extensão) e retorna (DataFrame, SLA de cada linha)."""
    if conteudo[:4] == b'PK\x03\x04':
        return _ler_xlsx(conteudo)                      # .xlsx
    if conteudo[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        return _ler_xls_binario(conteudo)               # .xls do Excel 97-2003
    for codificacao in ('utf-8-sig', 'utf-16', 'cp1252'):
        try:
            texto = conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
        inicio = texto[:5000].lower()
        if 'urn:schemas-microsoft-com:office:spreadsheet' in inicio and '<workbook' in inicio:
            return _ler_xls_xml(texto)                  # Planilha XML 2003 salva como .xls
        if '<table' in texto.lower():
            return _ler_xls_html(texto)                 # Página HTML salva como .xls
    raise ValueError("Formato de arquivo não reconhecido. Envie a planilha em .xls ou .xlsx.")


@st.cache_data(show_spinner="Lendo planilha...")
def carregar_dados(conteudo):
    df, status_sla_cores = ler_planilha(conteudo)
    faltando = [c for c in ['Estação', 'Status', 'Empresa', 'Data'] if c not in df.columns]
    if faltando:
        raise ValueError(f"a planilha não tem a(s) coluna(s) {', '.join(faltando)}. "
                         "Confira se é a exportação do Portal de Chamados.")

    status_sla_cores = status_sla_cores[:len(df)]
    if len(status_sla_cores) < len(df):
        status_sla_cores += ["SEM SLA"] * (len(df) - len(status_sla_cores))

    df['SLA'] = status_sla_cores

    # A planilha tem duas colunas "Data": abertura e acionamento (a 2ª vira "Data.1")
    df = df.rename(columns={'Data.1': 'Data Acionamento'})

    df = df.dropna(subset=['Estação'])
    # Chamado repetido (mesmo ID Sopro e mesma empresa) contaria em dobro: fica uma linha só, qualquer uma.
    # O mesmo ID com empresa diferente continua na contagem, uma vez por empresa.
    antes = len(df)
    if 'ID Sopro' in df.columns:
        sem_id = df['ID Sopro'].isna() | (df['ID Sopro'].astype(str).str.strip() == '')
        df = pd.concat([df[~sem_id].drop_duplicates(subset=['ID Sopro', 'Empresa']),
                        df[sem_id].drop_duplicates()]).sort_index()
    else:
        df = df.drop_duplicates()
    duplicadas = antes - len(df)
    # Colunas com números e textos misturados (ex.: "1213/ 1214") viram texto para exibição
    for col in ['ID Atividade', 'ID Fornecedor', 'Chamado']:
        if col in df.columns:
            df[col] = df[col].apply(lambda v: '' if pd.isna(v) else str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
    df['Status'] = df['Status'].astype(str).str.strip()
    df = df[~df['Status'].apply(status_excluido)]
    df['Empresa'] = df['Empresa'].fillna('NÃO INFORMADA').astype(str).str.strip()
    for col in ['UF', 'Regional']:
        if col in df.columns:
            df[col] = df[col].fillna('NÃO INFORMADA').astype(str).str.strip().str.upper()
    if 'Tipo' in df.columns:
        df['Tipo'] = df['Tipo'].fillna('NÃO INFORMADO').astype(str).str.strip()

    for col in ['Data', 'Data Acionamento', 'Data Fechamento']:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors='coerce')

    df['Categoria'] = df['Status'].apply(categorizar_status)

    if 'Data' in df.columns:
        df['Mês Abertura'] = df['Data'].dt.to_period('M').dt.to_timestamp()

    # Tempos em dias (valores negativos são inconsistências de digitação e são descartados)
    if 'Data' in df.columns and 'Data Fechamento' in df.columns:
        dias = (df['Data Fechamento'] - df['Data']).dt.days
        df['Dias p/ Resolver'] = dias.where((dias >= 0) & (df['Categoria'] == 'Resolvido'))
        df['Mês Fechamento'] = df['Data Fechamento'].dt.to_period('M').dt.to_timestamp()
    if 'Data' in df.columns and 'Data Acionamento' in df.columns:
        dias = (df['Data Acionamento'] - df['Data']).dt.days
        df['Dias p/ Acionar'] = dias.where(dias >= 0)

    return df, duplicadas


def preparar_filtro(key, opcoes, padrao):
    """Mantém a escolha do filtro entre execuções; só descarta o que saiu das opções atuais."""
    atual = st.session_state.get(key, padrao)
    if isinstance(padrao, list):
        atual = [v for v in atual if v in opcoes]
    elif atual not in opcoes:
        atual = padrao
    st.session_state[key] = atual


def data_da_extracao(nome_arquivo):
    """Data/hora da extração no nome do arquivo do portal (Chamados_AAAAMMDDhhmmss), ou None."""
    m = re.search(r'Chamados_(\d{14})', nome_arquivo or '')
    return pd.to_datetime(m.group(1), format='%Y%m%d%H%M%S', errors='coerce') if m else None


def formatar_data(valor, vazio='sem data'):
    """dd/mm/aaaa; data ausente (NaT/None) vira texto em vez de derrubar o app."""
    return valor.strftime('%d/%m/%Y') if pd.notna(valor) else vazio


def faixa_aging(dias):
    if pd.isna(dias):
        return "Sem data"
    if dias <= 30:
        return "0-30 dias"
    if dias <= 60:
        return "31-60 dias"
    if dias <= 90:
        return "61-90 dias"
    if dias <= 180:
        return "91-180 dias"
    return "+180 dias"


FAIXAS_AGING = ["0-30 dias", "31-60 dias", "61-90 dias", "91-180 dias", "+180 dias", "Sem data"]


def pct(parte, total):
    return (parte / total * 100) if total else 0.0


def resumo_por_grupo(df, grupo, data_ref):
    """Indicadores de desempenho agrupados por uma coluna (Empresa, Estação, Tipo...)."""
    g = df.groupby(grupo)
    r = pd.DataFrame({
        'Total': g.size(),
        'Resolvidos': g['Categoria'].apply(lambda x: (x == 'Resolvido').sum()),
        'Abertos': g['Categoria'].apply(lambda x: (x == 'Aberto').sum()),
        'Cancelados': g['Categoria'].apply(lambda x: (x == 'Cancelado').sum()),
        'SLA No Prazo': g['SLA'].apply(lambda x: (x == 'NO PRAZO').sum()),
        'SLA Fora do Prazo': g['SLA'].apply(lambda x: (x == 'FORA DO PRAZO').sum()),
    })
    validos = r['Total'] - r['Cancelados']
    r['% Resolução'] = (r['Resolvidos'] / validos.where(validos > 0) * 100).round(1)
    sla_medido = r['SLA No Prazo'] + r['SLA Fora do Prazo']
    r['% SLA No Prazo'] = (r['SLA No Prazo'] / sla_medido.where(sla_medido > 0) * 100).round(1)
    r['% Cancelados'] = (r['Cancelados'] / r['Total'] * 100).round(1)

    if 'Dias p/ Resolver' in df.columns:
        r['Média Dias Resolução'] = g['Dias p/ Resolver'].mean().round(1)
        r['Mediana Dias Resolução'] = g['Dias p/ Resolver'].median().round(1)
        r['P90 Dias Resolução'] = g['Dias p/ Resolver'].quantile(0.9).round(1)
    if 'Dias p/ Acionar' in df.columns:
        r['Mediana Dias Acionamento'] = g['Dias p/ Acionar'].median().round(1)

    abertos = df[df['Categoria'] == 'Aberto']
    if 'Data' in df.columns and not abertos.empty:
        idade = (data_ref - abertos['Data']).dt.days
        ab = abertos.assign(Idade=idade).groupby(grupo)['Idade']
        r['Idade Média Backlog (dias)'] = ab.mean().round(1)
        r['Backlog +90 dias'] = ab.apply(lambda x: (x > 90).sum())
    else:
        r['Idade Média Backlog (dias)'] = pd.NA
        r['Backlog +90 dias'] = 0
    r['Backlog +90 dias'] = r['Backlog +90 dias'].fillna(0).astype(int)

    return r.reset_index()


def calcular_score(r):
    """Índice de eficiência 0-100: 40% resolução, 30% SLA, 30% rapidez (mediana de dias, relativa)."""
    resol = r['% Resolução'].fillna(0) / 100
    sla = r['% SLA No Prazo'].fillna(0) / 100
    if 'Mediana Dias Resolução' in r.columns and r['Mediana Dias Resolução'].notna().any():
        med = r['Mediana Dias Resolução']
        mn, mx = med.min(), med.max()
        rapidez = (1 - (med - mn) / (mx - mn)) if mx > mn else pd.Series(1.0, index=r.index)
        rapidez = rapidez.fillna(0)
    else:
        rapidez = pd.Series(0.0, index=r.index)
    return ((0.4 * resol + 0.3 * sla + 0.3 * rapidez) * 100).round(1)


def historico_mensal(df, inicio=None, fim=None):
    """Abertos x Resolvidos x Cancelados por mês + backlog acumulado e tempo de resolução.

    inicio/fim (primeiro dia do mês) limitam o eixo aos meses do período escolhido na barra lateral.
    """
    abertos = df.groupby('Mês Abertura').size().rename('Abertos')
    res = df[df['Categoria'] == 'Resolvido']
    # Mês em que o chamado sai do backlog: fechamento; sem ele, acionamento; sem nenhum, abertura.
    # Sem isso, concluídos/cancelados sem data de fechamento ficariam no backlog para sempre.
    fechados = df[df['Categoria'].isin(['Resolvido', 'Cancelado'])]
    saida = fechados['Mês Fechamento']
    if 'Data Acionamento' in fechados.columns:
        saida = saida.fillna(fechados['Data Acionamento'].dt.to_period('M').dt.to_timestamp())
    saida = saida.fillna(fechados['Mês Abertura'])
    resolvidos = saida[fechados['Categoria'] == 'Resolvido'].value_counts().rename('Resolvidos')
    cancelados = saida[fechados['Categoria'] == 'Cancelado'].value_counts().rename('Cancelados')
    mediana = res.groupby('Mês Fechamento')['Dias p/ Resolver'].median().rename('Mediana Dias Resolução')
    media = res.groupby('Mês Fechamento')['Dias p/ Resolver'].mean().rename('Média Dias Resolução')

    h = pd.concat([abertos, resolvidos, cancelados, mediana, media], axis=1).sort_index()
    h.index.name = 'Mês'
    if h.empty:
        return h.reset_index()
    inicio = h.index.min() if inicio is None else inicio
    fim = h.index.max() if fim is None else fim
    h = h.reindex(pd.date_range(inicio, fim, freq='MS'))
    h.index.name = 'Mês'
    for c in ['Abertos', 'Resolvidos', 'Cancelados']:
        h[c] = h[c].fillna(0).astype(int)
    h['Saldo do Mês'] = h['Abertos'] - h['Resolvidos'] - h['Cancelados']
    h['Backlog Acumulado'] = h['Saldo do Mês'].cumsum()
    h['% Resolvidos s/ Abertos'] = (h['Resolvidos'] / h['Abertos'].where(h['Abertos'] > 0) * 100).round(1)
    h['Mediana Dias Resolução'] = h['Mediana Dias Resolução'].round(1)
    h['Média Dias Resolução'] = h['Média Dias Resolução'].round(1)
    return h.reset_index()


def filtrar_status(df, status_selecionado):
    if status_selecionado != "TODOS":
        return df[df['Status'] == status_selecionado]
    return df


def fechamento_antes_da_abertura(df):
    """Chamados com Data Fechamento anterior à Data de abertura (erro de digitação no portal)."""
    if 'Data Fechamento' not in df.columns:
        return pd.DataFrame()
    ruins = df[df['Data Fechamento'] < df['Data']].sort_values('Data Fechamento')
    t = ruins[[c for c in ['ID Sopro', 'Estação', 'Status', 'Empresa', 'Data', 'Data Fechamento'] if c in df.columns]].copy()
    for col in ['Data', 'Data Fechamento']:
        t[col] = t[col].map(formatar_data)
    return t.rename(columns={'Data': 'Aberto em', 'Data Fechamento': 'Fechado em'})


def comparar_trimestres(hist, data_ref):
    """Destaques dos últimos 3 meses completos contra os 3 anteriores."""
    if hist.empty:
        return []
    # O mês da data de referência só entra se já tiver terminado
    if (data_ref + pd.Timedelta(days=1)).month == data_ref.month:
        hist = hist[hist['Mês'] < data_ref.to_period('M').to_timestamp()]
    if len(hist) < 6:
        return []
    ult, ant = hist.tail(3), hist.iloc[-6:-3]
    variacao = lambda novo, base: f"{pct(novo - base, base):+.1f}%" if base else "—"
    meses = f"{ult['Mês'].iloc[0].strftime('%m/%Y')} a {ult['Mês'].iloc[-1].strftime('%m/%Y')}"
    ab_u, ab_a = ult['Abertos'].sum(), ant['Abertos'].sum()
    rs_u, rs_a = ult['Resolvidos'].sum(), ant['Resolvidos'].sum()
    destaques = [f"Nos últimos 3 meses completos ({meses}) foram abertos **{ab_u}** chamados "
                 f"({variacao(ab_u, ab_a)} vs. trimestre anterior) e resolvidos **{rs_u}** ({variacao(rs_u, rs_a)})."]
    med_u, med_a = ult['Mediana Dias Resolução'].mean(), ant['Mediana Dias Resolução'].mean()
    if pd.notna(med_u) and pd.notna(med_a):
        tend = "melhorou" if med_u < med_a else "piorou"
        destaques.append(f"O tempo mediano de resolução **{tend}**: {med_a:.0f} → {med_u:.0f} dias.")
    saldo = ult['Saldo do Mês'].sum()
    if saldo == 0:
        destaques.append("O backlog **ficou estável** no último trimestre.")
    else:
        destaques.append(f"O backlog {'cresceu' if saldo > 0 else 'diminuiu'} **{abs(saldo)}** chamados "
                         f"no último trimestre.")
    return destaques


def gerar_excel(abas):
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        for nome, tabela in abas.items():
            tabela.to_excel(writer, sheet_name=nome[:31], index=False)
            ws = writer.sheets[nome[:31]]
            ws.freeze_panes = 'A2'
            for col in ws.columns:
                largura = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(largura + 2, 10), 50)
    return buffer.getvalue()


def ranking_ocorrencias(df, filtro, titulo_periodo, escala, colunas_detalhe, opcoes_periodo, chave):
    if 'Data' not in df.columns or not pd.api.types.is_datetime64_any_dtype(df['Data']):
        st.error("Não foi possível ler as datas da planilha para gerar este relatório.")
        return

    periodo = st.radio("Selecione o período de análise:", list(opcoes_periodo), horizontal=True, key=chave)
    data_base = df['Data'].max()
    data_corte = data_base - pd.DateOffset(months=opcoes_periodo[periodo])

    df_sel = df[(df['Data'] >= data_corte) & filtro(df)]
    if df_sel.empty:
        st.success(f"Nenhum chamado de {titulo_periodo} encontrado no período selecionado.")
        return

    ranking = df_sel.groupby('Estação').agg(
        **{'Casos Registrados': ('Estação', 'size'),
           'Em Aberto': ('Categoria', lambda x: (x == 'Aberto').sum()),
           'Resolvidos': ('Categoria', lambda x: (x == 'Resolvido').sum()),
           'Último Registro': ('Data', 'max')}
    ).reset_index().sort_values('Casos Registrados', ascending=False)
    ranking['Último Registro'] = ranking['Último Registro'].map(formatar_data)

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Chamados no período", len(df_sel))
    k2.metric("Estações afetadas", ranking['Estação'].nunique())
    k3.metric("Reincidentes (2+ chamados)", int((ranking['Casos Registrados'] >= 2).sum()))
    k4.metric("Em aberto", int((df_sel['Categoria'] == 'Aberto').sum()))

    colA, colB = st.columns([1, 2])
    with colA:
        st.write(f"**Ranking (Desde {data_corte.strftime('%d/%m/%Y')} até {data_base.strftime('%d/%m/%Y')})**")
        st.dataframe(ranking, width='stretch', hide_index=True)
    with colB:
        top = ranking.head(30)
        fig = px.bar(top, x='Casos Registrados', y='Estação', orientation='h', text='Casos Registrados',
                     title=f'Top 30 estações ({periodo})', color='Casos Registrados', color_continuous_scale=escala)
        fig.update_layout(yaxis_categoryorder='total ascending', coloraxis_showscale=False,
                          height=altura_barras(len(top)))
        st.plotly_chart(fig, width='stretch')

    tend = df_sel[df_sel['Categoria'] == 'Aberto'].groupby('Mês Abertura').size().reset_index(name='Chamados')
    fig_t = px.bar(tend, x='Mês Abertura', y='Chamados', text='Chamados', title='Chamados abertos por mês',
                   color_discrete_sequence=[CATEGORIA_CORES['Aberto']])
    fig_t.update_xaxes(dtick='M1', tickformat='%m/%Y')
    st.plotly_chart(fig_t, width='stretch')

    with st.expander("🔎 Ver lista completa de chamados neste período"):
        cols = [c for c in colunas_detalhe if c in df_sel.columns]
        st.dataframe(df_sel[cols], width='stretch')


# ---------------------------------------------------------------------------
# Assistente de consulta por estação
# ---------------------------------------------------------------------------
COLUNAS_CONSULTA = ['ID Sopro', 'Data', 'Tipo', 'Descrição', 'Empresa', 'Status', 'Dias em aberto',
                    'Data Fechamento', 'Observações']
MAX_ESTACOES_RESPOSTA = 10

MENSAGEM_AJUDA = (
    "Digite o código de uma ou mais estações (ex.: **RSBGE02**) ou o ID de um chamado (ex.: **SPR00072339**). "
    "Também dá para digitar o início do código (ex.: **RSBGE**) para ver todas as estações daquele grupo."
)


def parece_codigo_estacao(token):
    """Letras e números juntos (ex.: RSBGE02), sem ser ID de chamado (SPR...)."""
    return bool(re.search(r'[A-Z]{2}', token) and re.search(r'\d', token)) and not re.match(r'SPR\d', token)


def identificar_estacoes(pergunta, estacoes):
    """Retorna (estações encontradas, sugestões, códigos digitados que não existem na planilha)."""
    mapa = {str(e).upper(): e for e in estacoes}
    tokens = [t.upper() for t in re.findall(r'[A-Za-z0-9]+', pergunta) if len(t) >= 4]

    encontradas = [mapa[t] for t in tokens if t in mapa]
    if encontradas:
        nao_encontrados = [t for t in dict.fromkeys(tokens) if t not in mapa and parece_codigo_estacao(t)]
        return list(dict.fromkeys(encontradas)), [], nao_encontrados

    # Início do código: "RSBGE" -> RSBGE02, RSBGE05...
    por_prefixo = [e for chave, e in sorted(mapa.items()) for t in tokens if chave.startswith(t)]
    if por_prefixo:
        return list(dict.fromkeys(por_prefixo)), [], []

    return [], sugestoes_parecidas(tokens, mapa), []


def sugestoes_parecidas(tokens, mapa):
    sugestoes = []
    for t in tokens:
        sugestoes += difflib.get_close_matches(t, list(mapa), n=5, cutoff=0.75)
    return [mapa[s] for s in dict.fromkeys(sugestoes)]


def tabela_chamados(chamados, data_ref):
    t = chamados.sort_values('Data', ascending=False).copy()
    t['Dias em aberto'] = (data_ref - t['Data']).dt.days.where(t['Categoria'] == 'Aberto').astype('Int64')
    for col in ['Data', 'Data Fechamento']:
        if col in t.columns:
            t[col] = t[col].map(lambda v: formatar_data(v, None if col == 'Data Fechamento' else 'sem data'))
    # Colunas vazias (ex.: fechamento de chamados abertos) não aparecem
    return t[[c for c in COLUNAS_CONSULTA if c in t.columns]].dropna(axis=1, how='all').reset_index(drop=True)


def responder_consulta(pergunta, df, data_ref):
    """Monta a resposta do assistente como uma lista de blocos {'texto': ..., 'tabela': DataFrame opcional}."""
    blocos = []

    # ID do chamado: "SPR00072339", sem os zeros ("SPR72339") ou só o número com 5+ dígitos ("72339")
    numeros = [int(n) for n in re.findall(r'\bSPR0*(\d+)\b', pergunta, flags=re.IGNORECASE)]
    numeros += [int(n) for n in re.findall(r'\b\d{5,}\b', pergunta)]
    if numeros and 'ID Sopro' in df.columns:
        num_df = pd.to_numeric(df['ID Sopro'].astype(str).str.extract(r'(\d+)')[0], errors='coerce')
        achados = df[num_df.isin(numeros)]
        existentes = set(num_df[num_df.isin(numeros)])
        faltando = [f"SPR{n:08d}" for n in dict.fromkeys(numeros) if n not in existentes]
        if faltando:
            blocos.append({'texto': f"Não encontrei o chamado **{', '.join(faltando)}** nesta planilha."})
        for _, c in achados.iterrows():
            fechamento = c.get('Data Fechamento')
            fech_txt = f" · fechado em {formatar_data(fechamento)}" if pd.notna(fechamento) else ""
            blocos.append({'texto': f"**{c['ID Sopro']}** · estação **{c['Estação']}** · {c.get('Tipo', '')} — "
                                    f"{c.get('Descrição', '')} · aberto em {formatar_data(c['Data'])} · "
                                    f"status **{c['Status']}** · {c.get('Empresa', '')}{fech_txt}"})
        if not re.sub(r'\bSPR\d+\b|\b\d{5,}\b', '', pergunta, flags=re.IGNORECASE).strip(' ,;e'):
            return blocos

    estacoes, sugestoes, nao_encontrados = identificar_estacoes(pergunta, df['Estação'].dropna().unique())
    # Códigos digitados junto com outros que existem: avisa em vez de ignorar em silêncio
    mapa = {str(e).upper(): e for e in df['Estação'].dropna().unique()}
    avisos = []
    for cod in nao_encontrados:
        parecidas = sugestoes_parecidas([cod], mapa)
        avisos.append({'texto': f"⚠️ Não encontrei a estação **{cod}** nesta planilha."
                                + (" Você quis dizer:" if parecidas else ""), 'estacoes': parecidas})
    if not estacoes:
        if blocos:
            return blocos
        texto = "Não reconheci nenhuma estação na sua mensagem."
        if sugestoes:
            return [{'texto': texto + " Você quis dizer:", 'estacoes': sugestoes}]
        return [{'texto': texto + "\n\n" + MENSAGEM_AJUDA}]

    if len(estacoes) > MAX_ESTACOES_RESPOSTA:
        abertos = df[df['Estação'].isin(estacoes) & (df['Categoria'] == 'Aberto')]
        resumo = abertos.groupby('Estação').agg(**{
            'Chamados em aberto': ('Estação', 'size'),
            'Mais antigo': ('Data', 'min'),
            'Tipos': ('Tipo', lambda x: ', '.join(sorted(x.unique()))),
        }).reset_index().sort_values('Chamados em aberto', ascending=False)
        resumo['Mais antigo'] = resumo['Mais antigo'].map(formatar_data)
        com_aberto = len(resumo)
        # Lista digitada código a código ou busca pelo início do código (ex.: RSBGE)
        digitadas = all(re.search(rf'\b{re.escape(str(e))}\b', pergunta, re.IGNORECASE) for e in estacoes)
        if digitadas:  # cada código digitado aparece, inclusive os sem chamado em aberto (🟢)
            sem_aberto = [e for e in estacoes if e not in set(resumo['Estação'])]
            resumo = pd.concat([resumo, pd.DataFrame({'Estação': sem_aberto, 'Chamados em aberto': 0,
                                                      'Mais antigo': '', 'Tipos': ''})], ignore_index=True)
        origem = "Você pesquisou" if digitadas else "Encontrei"
        complemento = "" if digitadas else " com esse início de código"
        blocos.append({'texto': f"{origem} **{len(estacoes)}** estações{complemento}; "
                                f"**{com_aberto}** delas têm chamado em aberto. Toque numa estação "
                                f"para ver os detalhes.",
                       'tabela': resumo if not resumo.empty else None})
        return blocos + avisos

    for est in estacoes:
        chamados = df[df['Estação'] == est]
        abertos = chamados[chamados['Categoria'] == 'Aberto']
        if not abertos.empty:
            tipos = ', '.join(sorted(abertos['Tipo'].unique())) if 'Tipo' in abertos else ''
            blocos.append({'texto': f"🔴 A estação **{est}** tem **{len(abertos)}** chamado(s) em aberto "
                                    f"({tipos}):",
                           'tabela': tabela_chamados(abertos, data_ref)})
        else:
            texto = f"🟢 A estação **{est}** não tem chamado em aberto."
            if chamados.empty:
                blocos.append({'texto': texto})
            else:
                blocos.append({'texto': texto + f" Chamados registrados ({len(chamados)}):",
                               'tabela': tabela_chamados(chamados, data_ref)})
    return blocos + avisos


# Botão flutuante, sempre visível no canto inferior direito, que abre o assistente em tela cheia
CSS_BOTAO_CONSULTA = """
<style>
/* 15px acima do padrão para não ficar atrás do botão do Streamlit no canto inferior direito */
.st-key-botao_consulta { position: fixed; right: 2rem; bottom: calc(2rem + 15px); z-index: 999; width: auto !important; }
.st-key-botao_consulta button { border-radius: 999px; padding: 1.1rem 2rem; min-height: 4rem; font-weight: 700;
  box-shadow: 0 6px 18px rgba(0, 0, 0, 0.25); }
.st-key-botao_consulta button p, .st-key-botao_consulta button span { font-size: 1.25rem; }
/* Espaço no fim da página para o último conteúdo não ficar embaixo do botão */
[data-testid="stMainBlockContainer"] { padding-bottom: 7rem; }
/* Celular: botão redondo só com o ícone; o texto continua para leitores de tela */
@media (max-width: 640px) { .st-key-botao_consulta { right: 1rem; bottom: calc(1rem + 15px); }
  .st-key-botao_consulta button { width: 3.75rem; height: 3.75rem; min-height: 0; padding: 0; justify-content: center; }
  .st-key-botao_consulta button span { font-size: 1.6rem; margin: 0; }
  .st-key-botao_consulta button p { position: absolute; width: 1px; height: 1px; overflow: hidden;
    clip: rect(0 0 0 0); white-space: nowrap; } }
</style>
"""

# Tela cheia do assistente: esconde barra lateral, cabeçalho e painel; deixa só a seta de voltar e o chat
CSS_TELA_CONSULTA = """
<style>
[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"], [data-testid="stHeader"],
.st-key-cabecalho_painel { display: none !important; }
[data-testid="stMainBlockContainer"] { padding-top: 4.5rem !important; }
.st-key-topo_consulta { position: fixed; top: 0; left: 0; right: 0; width: auto !important; z-index: 990;
  padding: 0.5rem 1rem; background: var(--st-background-color, #ffffff);
  border-bottom: 1px solid rgba(128, 128, 128, 0.2); }
.st-key-voltar_consulta button { border: none; background: transparent; padding: 0.1rem 0.6rem;
  min-height: 44px; min-width: 44px; }
[data-testid="stChatInputSubmitButton"] { min-height: 44px; min-width: 44px; }
.st-key-voltar_consulta button p { font-size: 2rem; line-height: 1; }
.st-key-topo_consulta h3 { margin: 0; padding: 0; }
[data-testid="stChatMessage"] { scroll-margin-top: 4.5rem; }
/* Fonte de 16px no campo de texto: abaixo disso o iPhone dá zoom na página ao focar no campo */
[data-testid="stChatInput"] textarea { font-size: 16px !important; }
</style>
"""

# Ajustes para o teclado do celular:
# - Android: com interactive-widget=resizes-content, o teclado encolhe a página e o campo fica acima dele
# - iPhone: acompanha a área visível (visualViewport) e sobe o campo de texto a altura do teclado
# - depois de cada resposta, rola até a última mensagem
JS_TECLADO = """
export default function(component) {
  const meta = document.querySelector('meta[name="viewport"]');
  if (meta && !meta.content.includes('interactive-widget')) {
    meta.content += ', interactive-widget=resizes-content';
  }
  const vv = window.visualViewport;
  const ajustar = () => {
    const fundo = document.querySelector('[data-testid="stBottom"]');
    if (!fundo || !vv) return;
    const teclado = Math.max(0, window.innerHeight - vv.height - vv.offsetTop);
    fundo.style.transform = teclado > 0 ? 'translateY(-' + teclado + 'px)' : '';
  };
  if (vv && !window.__consultaTeclado) {
    window.__consultaTeclado = true;
    vv.addEventListener('resize', ajustar);
    vv.addEventListener('scroll', ajustar);
  }
  ajustar();
  // Mostra a última pergunta no topo da tela, com a resposta logo abaixo, para ler desde o começo.
  // Repete por pouco mais de 1 s porque o Streamlit rola sozinho para o fim enquanto as tabelas carregam.
  const mostrarUltimaPergunta = () => {
    const msgs = document.querySelectorAll('[data-testid="stChatMessage"]');
    if (msgs.length) msgs[Math.max(0, msgs.length - 2)].scrollIntoView({ block: 'start' });
  };
  [150, 500, 900, 1400].forEach((ms) => setTimeout(mostrarUltimaPergunta, ms));
}
"""

_ajuste_teclado = st.components.v2.component("ajuste_teclado_consulta", js=JS_TECLADO)

# Barra do topo: o Streamlit Cloud adiciona botões (compartilhar, estrela, editar, GitHub...).
# Esconde todos, menos o do GitHub. O ícone de cada botão é uma imagem de fundo, então a escolha
# é feita pelo endereço da imagem; se nenhum for reconhecido como GitHub, nada é escondido.
JS_BARRA_TOPO = """
export default function() {
  if (window.__barraTopoSopro) return;
  window.__barraTopoSopro = true;
  const ehGithub = (botao) => {
    const icone = botao.querySelector('[data-testid="stToolbarActionButtonIcon"]');
    const fundo = icone ? getComputedStyle(icone).backgroundImage : '';
    let texto = fundo + ' ' + (botao.textContent || '');
    const base64 = fundo.match(/base64,([^"')]+)/);
    if (base64) { try { texto += ' ' + atob(base64[1]); } catch (e) {} }
    try { texto += ' ' + decodeURIComponent(fundo); } catch (e) {}
    return /github/i.test(texto);
  };
  let agendado = false;
  const aplicar = () => {
    agendado = false;
    const botoes = [...document.querySelectorAll('[data-testid="stToolbarActionButton"]')];
    const github = botoes.filter(ehGithub);
    if (!github.length) return;
    botoes.forEach((b) => { b.style.display = github.includes(b) ? '' : 'none'; });
  };
  new MutationObserver(() => {
    if (!agendado) { agendado = true; requestAnimationFrame(aplicar); }
  }).observe(document.body, { childList: true, subtree: true });
  aplicar();
}
"""

_ajuste_barra_topo = st.components.v2.component("ajuste_barra_topo", js=JS_BARRA_TOPO)


@st.fragment
def botao_consulta():
    # Fragmento: o clique só marca o modo de consulta e recarrega a página já na tela do chat
    st.html(CSS_BOTAO_CONSULTA)
    if st.button("Consultar estação", icon="💬", key="botao_consulta", type="primary",
                 help="Ver se uma estação já tem chamado aberto"):
        st.session_state.modo_consulta = True
        st.rerun()


CORES_CATEGORIA = {'Aberto': 'orange', 'Resolvido': 'green', 'Cancelado': 'gray'}


def botao_pesquisar(estacao, key):
    st.button(f"🔎 {estacao}", key=key, on_click=st.session_state.__setitem__,
              args=('consulta_pendente', estacao))


def mostrar_cartoes(tabela, key):
    """Um cartão por linha: no celular cabe sem arrastar para o lado, ao contrário de uma tabela larga."""
    for n, linha in tabela.iterrows():
        with st.container(border=True):
            if 'ID Sopro' in linha:
                st.markdown(f"**{linha['ID Sopro']}** · {linha.get('Tipo', '')}")
                st.badge(linha['Status'], color=CORES_CATEGORIA[categorizar_status(linha['Status'])])
                st.markdown(str(linha.get('Descrição', '')))
                detalhes = [f"Aberto em {linha['Data']}"]
                if pd.notna(linha.get('Dias em aberto')):
                    detalhes.append(f"**{linha['Dias em aberto']} dias em aberto**")
                if pd.notna(linha.get('Data Fechamento')):
                    detalhes.append(f"fechado em {linha['Data Fechamento']}")
                detalhes.append(str(linha.get('Empresa', '')))
                st.caption(' · '.join(detalhes))
                obs = linha.get('Observações')
                if pd.notna(obs) and str(obs).strip(' -'):
                    with st.expander("Ver observações"):
                        st.write(obs)
            else:  # resumo de um grupo de estações (busca pelo início do código)
                botao_pesquisar(linha['Estação'], f'grp_{key}_{n}')
                if linha['Chamados em aberto']:
                    st.caption(f"🔴 **{linha['Chamados em aberto']} em aberto** · mais antigo: "
                               f"{linha['Mais antigo']} · {linha['Tipos']}")
                else:
                    st.caption("🟢 sem chamado em aberto")


def tela_consulta(df, data_ref):
    """Assistente em tela cheia, com o campo de texto preso no rodapé (fica acima do teclado no celular)."""
    st.html(CSS_TELA_CONSULTA)

    with st.container(key="topo_consulta", horizontal=True, vertical_alignment="center"):
        if st.button("←", key="voltar_consulta", help="Voltar ao painel"):
            st.session_state.modo_consulta = False
            st.rerun()
        st.subheader("Consultar Estação")

    if 'chat_consulta' not in st.session_state:
        st.session_state.chat_consulta = [{'papel': 'assistant', 'blocos': [{'texto': "Olá! " + MENSAGEM_AJUDA}]}]

    # Tocar numa estação sugerida (ou no resumo de um grupo) pesquisa por ela, sem digitar
    pergunta = st.chat_input("Estação ou ID do chamado") or st.session_state.pop('consulta_pendente', None)
    if pergunta:
        st.session_state.chat_consulta.append({'papel': 'user', 'blocos': [{'texto': pergunta}]})
        st.session_state.chat_consulta.append({'papel': 'assistant',
                                               'blocos': responder_consulta(pergunta, df, data_ref)})

    for i, msg in enumerate(st.session_state.chat_consulta):
        with st.chat_message(msg['papel']):
            for j, bloco in enumerate(msg['blocos']):
                st.markdown(bloco['texto'])
                if bloco.get('tabela') is not None:
                    mostrar_cartoes(bloco['tabela'], f'{i}_{j}')
                for est in bloco.get('estacoes', []):
                    botao_pesquisar(est, f'sug_{i}_{j}_{est}')

    if len(st.session_state.chat_consulta) > 1 and st.button("🗑️ Limpar conversa"):
        del st.session_state.chat_consulta
        st.rerun()

    _ajuste_teclado(key="ajuste_teclado", data={'mensagens': len(st.session_state.chat_consulta)})


# ---------------------------------------------------------------------------
# Seletor de mês/ano: duas listas suspensas lado a lado
# ---------------------------------------------------------------------------
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]


def seletor_mes_ano(titulo, key, valor_inicial, ano_min, ano_max):
    """Listas suspensas de mês e ano lado a lado. Retorna o primeiro dia do mês escolhido."""
    anos = list(range(ano_min, ano_max + 1))
    st.caption(titulo)
    col_mes, col_ano = st.columns([3, 2])
    mes = col_mes.selectbox(f"Mês ({titulo})", MESES, index=valor_inicial.month - 1,
                            key=f'{key}_mes', label_visibility="collapsed", placeholder="Selecione")
    ano = col_ano.selectbox(f"Ano ({titulo})", anos, index=anos.index(valor_inicial.year),
                            key=f'{key}_ano', label_visibility="collapsed", placeholder="Selecione")
    return pd.Timestamp(year=ano, month=MESES.index(mes) + 1, day=1)


# ---------------------------------------------------------------------------
# Interface
# ---------------------------------------------------------------------------
modo_consulta = st.session_state.get('modo_consulta', False)

st.html("<style>.st-key-ajuste_barra_topo { display: none; }"
        # No celular, título menor para o envio do arquivo aparecer sem rolar
        "@media (max-width: 640px) { h1 { font-size: 1.9rem !important; } }"
        # Botões com área de toque confortável (mínimo de 44 px)
        "[data-testid=stBaseButton-secondary] { min-height: 44px; }"
        # Nomes longos (ex.: empresa mais eficiente) quebram linha em vez de ficarem cortados
        "[data-testid=stMetricValue] p { white-space: normal !important; overflow-wrap: anywhere; }"
        # No celular, a barra de ferramentas dos gráficos só atrapalha o toque
        "@media (max-width: 640px) { .modebar-container { display: none !important; } }</style>")
_ajuste_barra_topo(key="ajuste_barra_topo")

with st.container(key="cabecalho_painel"):
    st.title("📊 Gestão de Chamados SOPRO")
    st.write("Suba aqui o arquivo extraído do portal, em .xls ou .xlsx")
    arquivo_upload = st.file_uploader("Suba o arquivo Excel aqui", type=["xlsx", "xls"])

if arquivo_upload:
    try:
        df, duplicadas = carregar_dados(arquivo_upload.getvalue())
    except Exception as erro:
        st.error(f"Não consegui ler o arquivo **{arquivo_upload.name}**: {erro}")
        st.stop()
    if duplicadas:
        st.warning(f"{duplicadas} linha(s) repetida(s) na planilha foram descartadas (mesmo ID Sopro e mesma empresa de outra linha).")
    desconhecidos = status_desconhecidos(df)
    if desconhecidos:
        st.warning(f"Status não reconhecido(s): {', '.join(desconhecidos)}. Foram contados como **em aberto**; "
                   "avise para incluí-los na lista de status.")
    # A conversa do chat vale para o arquivo carregado: arquivo diferente, conversa nova
    impressao = hashlib.md5(arquivo_upload.getvalue()).hexdigest()
    if st.session_state.get('arquivo_do_chat') != impressao:
        st.session_state.pop('chat_consulta', None)
        st.session_state['arquivo_do_chat'] = impressao
    # "Dias em aberto" e idade do backlog contam até a extração do arquivo; sem a data no nome, até hoje
    extracao = data_da_extracao(arquivo_upload.name)
    data_ref = (extracao if pd.notna(extracao) else pd.Timestamp.today()).normalize()

    if not modo_consulta:
        botao_consulta()

    st.sidebar.header("🎯 Funil de Filtros")

    periodo_ini = periodo_fim = None
    if 'Data' in df.columns and df['Data'].notna().any():
        mes_min = df['Data'].min().to_period('M').to_timestamp()
        mes_max = df['Data'].max().to_period('M').to_timestamp()
        versao = st.session_state.get('periodo_versao', 0)

        st.sidebar.markdown("**0. Período de abertura:**")
        with st.sidebar:
            periodo_ini = seletor_mes_ano("De", f'periodo_ini_{versao}', mes_min, mes_min.year, mes_max.year)
            periodo_fim = seletor_mes_ano("Até", f'periodo_fim_{versao}', mes_max, mes_min.year, mes_max.year)
        if st.sidebar.button("↺ Todo o período"):
            st.session_state['periodo_versao'] = versao + 1
            st.rerun()

        if periodo_ini > periodo_fim:
            periodo_ini, periodo_fim = periodo_fim, periodo_ini
            st.sidebar.caption("O mês inicial era posterior ao final; os dois foram invertidos.")
        periodo_ini, periodo_fim = max(periodo_ini, mes_min), min(periodo_fim, mes_max)
        st.sidebar.caption(f"Considerando de {periodo_ini.strftime('%m/%Y')} a {periodo_fim.strftime('%m/%Y')}.")

        df_f0 = df[(df['Data'] >= periodo_ini) & (df['Data'] < periodo_fim + pd.DateOffset(months=1))]
    else:
        df_f0 = df

    # Regional = estado. A coluna UF traz o estado; se a planilha não tiver UF, usa a coluna Regional.
    # Filtrar aqui (antes da empresa) faz todas as abas, inclusive Eficiência, Ninhos e Zeladorias,
    # considerarem só a regional escolhida.
    col_regional = next((c for c in ['UF', 'Regional'] if c in df.columns), None)
    if col_regional:
        regionais = sorted(df[col_regional].unique().tolist())
        preparar_filtro('f_regional', ["TODAS"] + regionais, "TODAS")
        regional_selecionada = st.sidebar.selectbox("1. Regional (Estado):", options=["TODAS"] + regionais,
                                                    key='f_regional', placeholder="Selecione")
        if regional_selecionada != "TODAS":
            df_f0 = df_f0[df_f0[col_regional] == regional_selecionada]

    empresas_disponiveis = sorted(df_f0['Empresa'].dropna().astype(str).unique().tolist())
    preparar_filtro('f_empresa', ["TODAS"] + empresas_disponiveis, "TODAS")
    empresa_selecionada = st.sidebar.selectbox("2. Empresa Prestadora:", options=["TODAS"] + empresas_disponiveis,
                                               key='f_empresa', placeholder="Selecione")

    if empresa_selecionada != "TODAS":
        df_f1 = df_f0[df_f0['Empresa'] == empresa_selecionada]
    else:
        df_f1 = df_f0

    opcoes_status = ["TODOS"] + sorted(df_f1['Status'].unique().tolist())
    preparar_filtro('f_status', opcoes_status, "TODOS")
    status_selecionado = st.sidebar.selectbox("3. Status do Chamado:", options=opcoes_status, key='f_status',
                                              placeholder="Selecione")

    df_f2 = filtrar_status(df_f1, status_selecionado)

    sla_disponiveis = sorted(df_f2['SLA'].dropna().astype(str).unique().tolist())
    preparar_filtro('f_sla', ["TODOS"] + sla_disponiveis, "TODOS")
    sla_selecionado = st.sidebar.selectbox("4. SLA:", options=["TODOS"] + sla_disponiveis, key='f_sla',
                                           placeholder="Selecione")

    if sla_selecionado != "TODOS":
        df_f3 = df_f2[df_f2['SLA'].astype(str) == sla_selecionado]
    else:
        df_f3 = df_f2

    if 'Tipo' in df_f3.columns:
        tipos_disponiveis = sorted(df['Tipo'].unique().tolist())
        preparar_filtro('f_tipos', tipos_disponiveis, tipos_disponiveis)
        tipos_selecionados = st.sidebar.multiselect("5. Tipo de Chamado:", options=tipos_disponiveis, key='f_tipos',
                                                    placeholder="Selecione")
        df_f4 = df_f3[df_f3['Tipo'].isin(tipos_selecionados)]
    else:
        df_f4 = df_f3

    estacoes_disponiveis = sorted(df_f4['Estação'].dropna().astype(str).unique().tolist())
    preparar_filtro('f_estacoes', estacoes_disponiveis, [])
    estacoes_selecionadas = st.sidebar.multiselect(
        "6. Estações (Sites):", options=estacoes_disponiveis, key='f_estacoes', placeholder="Todas (toque para escolher)",
        help="Deixe vazio para considerar todas as estações."
    )

    df_final = df_f4[df_f4['Estação'].isin(estacoes_selecionadas)] if estacoes_selecionadas else df_f4

    # O assistente é desenhado depois dos filtros (escondidos), para que eles não percam a seleção
    if modo_consulta:
        tela_consulta(df, data_ref)
        st.stop()

    if df_final.empty:
        st.warning("Nenhum chamado encontrado com os filtros selecionados.")
        st.stop()

    # No celular a barra lateral começa fechada e só aparece o ícone »: avisa onde ficam os filtros
    st.html("<style>.dica-filtros{display:none} @media (max-width:640px){.dica-filtros{display:block;"
            "padding:.6rem .8rem;border-radius:.5rem;background:#EEF4FF;color:#1F3A68;font-size:1rem}}</style>"
            "<div class='dica-filtros'>🎯 <b>Filtros</b>: toque em <b>»</b> no canto superior esquerdo "
            "para escolher período, estado, empresa e status.</div>")
    referencia = (f"Dados extraídos em **{extracao.strftime('%d/%m/%Y %H:%M')}**" if pd.notna(extracao) else
                  f"Dias em aberto contados até hoje, **{data_ref.strftime('%d/%m/%Y')}** "
                  f"(o nome do arquivo não traz a data da extração)")
    st.caption(f"{referencia} · {len(df_final)} chamados no filtro atual")

    aba1, aba2, aba3, aba4, aba5, aba6 = st.tabs([
        "📊 Visão Geral", "📈 Histórico de Resolução", "🏢 Eficiência por Empresa",
        "🧾 Resumo Executivo", "🐦 Ninhos na EV", "🌿 Zeladorias"
    ])

    # ----------------------------------------------------------------- Visão Geral
    with aba1:
        st.subheader("📈 Resumo do Filtro Atual")
        total_chamados = len(df_final)
        cat = df_final['Categoria'].value_counts()
        qtd_aberto = int(cat.get('Aberto', 0))
        qtd_resolvido = int(cat.get('Resolvido', 0))
        qtd_cancelado = int(cat.get('Cancelado', 0))

        qtd_no_prazo = int((df_final['SLA'] == 'NO PRAZO').sum())
        qtd_sla_medido = int(df_final['SLA'].isin(['NO PRAZO', 'FORA DO PRAZO']).sum())
        mediana_res = df_final['Dias p/ Resolver'].median() if 'Dias p/ Resolver' in df_final else None

        m1, m2, m3, m4, m5, m6 = st.columns(6)
        m1.metric("Total de Chamados", total_chamados)
        m2.metric("% Em Aberto", f"{pct(qtd_aberto, total_chamados):.1f}%", f"{qtd_aberto} chamados", delta_color="off")
        m3.metric("% Resolvidos", f"{pct(qtd_resolvido, total_chamados):.1f}%", f"{qtd_resolvido} chamados", delta_color="off")
        m4.metric("% Cancelados", f"{pct(qtd_cancelado, total_chamados):.1f}%", f"{qtd_cancelado} chamados", delta_color="off")
        m5.metric("% SLA (No Prazo)", f"{pct(qtd_no_prazo, qtd_sla_medido):.1f}%",
                  help="Chamados com bolinha verde ÷ chamados com bolinha verde ou vermelha.")
        m6.metric("Mediana p/ Resolver", f"{mediana_res:.0f} dias" if pd.notna(mediana_res) else "—",
                  help="Dias entre a abertura e a data de fechamento dos chamados concluídos/fechados.")

        st.divider()

        c1, c2 = st.columns(2)
        with c1:
            st.subheader("Status dos chamados")
            df_status = df_final.groupby(['Categoria', 'Status']).size().reset_index(name='Quantidade')
            fig_sun = px.sunburst(df_status, path=['Categoria', 'Status'], values='Quantidade',
                                  color='Categoria', color_discrete_map=CATEGORIA_CORES)
            fig_sun.update_traces(textinfo='label+percent root')
            st.plotly_chart(fig_sun, width='stretch')
        with c2:
            st.subheader("SLA por tipo de chamado")
            if 'Tipo' in df_final.columns:
                df_sla_tipo = df_final.groupby(['Tipo', 'SLA']).size().reset_index(name='Quantidade')
                fig_sla = px.bar(df_sla_tipo, x='Tipo', y='Quantidade', color='SLA', barmode='stack',
                                 color_discrete_map=SLA_CORES)
                st.plotly_chart(fig_sla, width='stretch')

        st.subheader("📊 Distribuição de Status por Estação (Top 30 em volume)")
        top_estacoes = df_final['Estação'].value_counts().head(30).index
        df_grafico = df_final[df_final['Estação'].isin(top_estacoes)].groupby(['Estação', 'Categoria']).size().reset_index(name='Quantidade')
        df_grafico['Total_Estacao'] = df_grafico.groupby('Estação')['Quantidade'].transform('sum')
        df_grafico['Porcentagem'] = (df_grafico['Quantidade'] / df_grafico['Total_Estacao'] * 100).round(1)

        fig = px.bar(
            df_grafico, x='Quantidade', y='Estação', color='Categoria', orientation='h',
            text='Porcentagem', barmode='stack', color_discrete_map=CATEGORIA_CORES
        )
        fig.update_traces(texttemplate='%{text}%', textposition='inside')
        fig.update_layout(yaxis_categoryorder='total ascending', height=altura_barras(len(top_estacoes)))
        st.plotly_chart(fig, width='stretch')

        st.divider()

        st.subheader("🏢 Relatório por Estação")
        relatorio_estacao = resumo_por_grupo(df_final, 'Estação', data_ref).sort_values('Total', ascending=False)
        st.dataframe(relatorio_estacao, width='stretch', hide_index=True)

        csv_relatorio = relatorio_estacao.to_csv(index=False, sep=';').encode('utf-8-sig')
        st.download_button(label="📥 Baixar Relatório (CSV)", data=csv_relatorio, file_name='relatorio.csv', mime='text/csv')

        with st.expander("🔎 Ver dados detalhados deste filtro"):
            st.dataframe(df_final, width='stretch')

            csv_detalhado = df_final.to_csv(index=False, sep=';').encode('utf-8-sig')
            nome_arquivo = re.sub(r'[^\w-]+', '_', empresa_selecionada).strip('_')  # "OI S/A" -> "OI_S_A"
            st.download_button(
                label=f"Baixar Dados Detalhados - {empresa_selecionada} (CSV)",
                data=csv_detalhado,
                file_name=f"dados_detalhados_{nome_arquivo}.csv",
                mime='text/csv',
            )

    # ------------------------------------------------------- Histórico de Resolução
    with aba2:
        st.subheader("📈 Histórico de Resolução")
        st.write("Chamados abertos por mês (data de abertura) comparados com os resolvidos e cancelados "
                 "(data de fechamento). O backlog acumulado mostra se a fila está crescendo ou diminuindo.")

        hist = historico_mensal(df_final, periodo_ini, periodo_fim)
        if hist.empty:
            st.info("Sem dados suficientes para montar o histórico.")
        else:
            fig_h = make_subplots(specs=[[{"secondary_y": True}]])
            fig_h.add_bar(x=hist['Mês'], y=hist['Abertos'], name='Abertos', marker_color='#3498DB')
            fig_h.add_bar(x=hist['Mês'], y=hist['Resolvidos'], name='Resolvidos', marker_color=CATEGORIA_CORES['Resolvido'])
            fig_h.add_bar(x=hist['Mês'], y=hist['Cancelados'], name='Cancelados', marker_color=CATEGORIA_CORES['Cancelado'])
            fig_h.add_trace(go.Scatter(x=hist['Mês'], y=hist['Backlog Acumulado'], name='Backlog acumulado',
                                       mode='lines+markers', line=dict(color='#E74C3C', width=3)), secondary_y=True)
            fig_h.update_layout(barmode='group', title='Abertos x Resolvidos x Cancelados por mês',
                                hovermode='x unified')
            fig_h.update_xaxes(tickformat='%m/%Y')
            fig_h.update_yaxes(title_text='Chamados no mês', secondary_y=False)
            fig_h.update_yaxes(title_text='Backlog acumulado', secondary_y=True)
            st.plotly_chart(fig_h, width='stretch')

            c1, c2 = st.columns(2)
            with c1:
                fig_tempo = go.Figure()
                fig_tempo.add_trace(go.Scatter(x=hist['Mês'], y=hist['Mediana Dias Resolução'], name='Mediana',
                                               mode='lines+markers', line=dict(color='#8E44AD', width=3)))
                fig_tempo.add_trace(go.Scatter(x=hist['Mês'], y=hist['Média Dias Resolução'], name='Média',
                                               mode='lines', line=dict(color='#8E44AD', dash='dot')))
                fig_tempo.update_layout(title='Dias para resolver, por mês',
                                        hovermode='x unified')
                fig_tempo.update_xaxes(tickformat='%m/%Y')
                st.plotly_chart(fig_tempo, width='stretch')
            with c2:
                fig_taxa = px.line(hist, x='Mês', y='% Resolvidos s/ Abertos', markers=True,
                                   title='Resolvidos ÷ abertos no mês (%)')
                fig_taxa.add_hline(y=100, line_dash='dash', line_color='gray',
                                   annotation_text='100% = fila estável')
                fig_taxa.update_xaxes(tickformat='%m/%Y')
                st.plotly_chart(fig_taxa, width='stretch')

            # Tempo de resolução por coorte de abertura e tipo
            res = df_final[df_final['Categoria'] == 'Resolvido'].dropna(subset=['Dias p/ Resolver'])
            if not res.empty and 'Tipo' in res.columns:
                c3, c4 = st.columns(2)
                with c3:
                    fig_box = px.box(res, x='Tipo', y='Dias p/ Resolver', color='Tipo', points=False,
                                     title='Dias para resolver, por tipo')
                    fig_box.update_layout(showlegend=False)
                    st.plotly_chart(fig_box, width='stretch')
                with c4:
                    curva = res['Dias p/ Resolver'].sort_values().reset_index(drop=True)
                    acum = pd.DataFrame({'Dias': curva, '% Resolvidos': (curva.index + 1) / len(curva) * 100})
                    # Percentuais no título: como rótulos nas linhas, 30/60/90 dias ficavam um em cima do outro
                    marcos = ' · '.join(f'{d}d: {(curva <= d).mean() * 100:.0f}%' for d in [30, 60, 90])
                    fig_cdf = px.line(acum, x='Dias', y='% Resolvidos',
                                      title=f'% resolvidos até X dias<br><sup>{marcos}</sup>')
                    for d in [30, 60, 90]:
                        fig_cdf.add_vline(x=d, line_dash='dot', line_color='gray')
                    st.plotly_chart(fig_cdf, width='stretch')

            # Backlog atual por idade
            abertos = df_final[df_final['Categoria'] == 'Aberto'].copy()
            if not abertos.empty:
                st.subheader("⏳ Idade do backlog em aberto")
                abertos['Idade (dias)'] = (data_ref - abertos['Data']).dt.days
                abertos['Faixa'] = abertos['Idade (dias)'].apply(faixa_aging)
                aging = abertos.groupby(['Faixa', 'Tipo' if 'Tipo' in abertos else 'Status']).size().reset_index(name='Chamados')
                fig_aging = px.bar(aging, x='Faixa', y='Chamados', color=aging.columns[1], barmode='stack',
                                   category_orders={'Faixa': FAIXAS_AGING},
                                   title=f'{len(abertos)} em aberto, por idade')
                st.plotly_chart(fig_aging, width='stretch')

            with st.expander("📋 Tabela do histórico mensal"):
                hist_tab = hist.copy()
                hist_tab['Mês'] = hist_tab['Mês'].dt.strftime('%m/%Y')
                st.dataframe(hist_tab, width='stretch', hide_index=True)

        incons = fechamento_antes_da_abertura(df_final)
        if not incons.empty:
            with st.expander(f"⚠️ Inconsistências da planilha: {len(incons)} chamado(s) fechado(s) antes da abertura"):
                st.caption("Ficam fora do tempo de resolução, mas contam como fechados no mês da data de fechamento.")
                st.dataframe(incons, width='stretch', hide_index=True)

    # ------------------------------------------------------- Eficiência por Empresa
    with aba3:
        st.subheader("🏢 Eficiência das Empresas Prestadoras")
        st.write("O **índice de eficiência** (0-100) combina: 40% taxa de resolução (resolvidos ÷ chamados não "
                 "cancelados), 30% SLA no prazo e 30% rapidez (mediana de dias para resolver, comparada entre as "
                 "empresas exibidas).")

        # A comparação usa período, tipo e estações da barra lateral, mas ignora os filtros de
        # empresa, status e SLA, que distorceriam as taxas de resolução e de SLA.
        df_emp_base = df_f0
        if 'Tipo' in df_emp_base.columns:
            df_emp_base = df_emp_base[df_emp_base['Tipo'].isin(tipos_selecionados)]
        if estacoes_selecionadas:
            df_emp_base = df_emp_base[df_emp_base['Estação'].isin(estacoes_selecionadas)]
        if empresa_selecionada != "TODAS" or status_selecionado != "TODOS" or sla_selecionado != "TODOS":
            st.info("Os filtros de empresa, status e SLA não se aplicam a esta aba, para que todas as "
                    "empresas sejam comparadas sobre a mesma base de chamados.")

        vol_min = st.slider("Volume mínimo de chamados para entrar no ranking:", 1, 200, 20, step=1)
        emp = resumo_por_grupo(df_emp_base, 'Empresa', data_ref)
        emp = emp[emp['Total'] >= vol_min].copy()

        if emp.empty:
            st.info("Nenhuma empresa atinge o volume mínimo selecionado.")
        else:
            emp['Índice de Eficiência'] = calcular_score(emp)
            emp = emp.sort_values('Índice de Eficiência', ascending=False)
            ordem = emp['Empresa'].tolist()

            melhor, pior = emp.iloc[0], emp.iloc[-1]
            k1, k2, k3 = st.columns(3)
            k1.metric("Empresas no ranking", len(emp))
            k2.metric("🥇 Mais eficiente", melhor['Empresa'], f"índice {melhor['Índice de Eficiência']:.1f}",
                      delta_color="off")
            k3.metric("⚠️ Menos eficiente", pior['Empresa'], f"índice {pior['Índice de Eficiência']:.1f}",
                      delta_color="off")

            fig_score = px.bar(emp, x='Índice de Eficiência', y='Empresa', orientation='h',
                               color='Índice de Eficiência', color_continuous_scale='RdYlGn', range_color=[0, 100],
                               text='Índice de Eficiência', title='Índice de eficiência por empresa')
            fig_score.update_layout(yaxis={'categoryorder': 'total ascending'}, height=altura_barras(len(emp)),
                                    coloraxis_showscale=False)
            encurtar_nomes(fig_score, emp['Empresa'])
            st.plotly_chart(fig_score, width='stretch')

            c1, c2 = st.columns(2)
            with c1:
                comp = emp.melt(id_vars='Empresa', value_vars=['% Resolução', '% SLA No Prazo'],
                                var_name='Indicador', value_name='%')
                fig_comp = px.bar(comp, x='%', y='Empresa', color='Indicador', barmode='group', orientation='h',
                                  category_orders={'Empresa': ordem[::-1]}, title='Taxa de resolução x SLA no prazo',
                                  color_discrete_map={'% Resolução': '#2ECC71', '% SLA No Prazo': '#3498DB'})
                fig_comp.update_layout(height=altura_barras(2 * len(ordem)))
                encurtar_nomes(fig_comp, ordem)
                st.plotly_chart(fig_comp, width='stretch')
            with c2:
                fig_disp = px.scatter(emp, x='Mediana Dias Resolução', y='% Resolução', size='Total',
                                      color='% SLA No Prazo', color_continuous_scale='RdYlGn', hover_name='Empresa',
                                      title='Rapidez × resolução (bolha = volume)', size_max=60)
                st.plotly_chart(fig_disp, width='stretch')

            c3, c4 = st.columns(2)
            with c3:
                mix = df_emp_base[df_emp_base['Empresa'].isin(ordem)].groupby(['Empresa', 'Categoria']).size().reset_index(name='Chamados')
                fig_mix = px.bar(mix, x='Chamados', y='Empresa', color='Categoria', barmode='stack', orientation='h',
                                 color_discrete_map=CATEGORIA_CORES, category_orders={'Empresa': ordem[::-1]},
                                 title='Situação dos chamados por empresa')
                fig_mix.update_layout(height=altura_barras(len(ordem)))
                encurtar_nomes(fig_mix, ordem)
                st.plotly_chart(fig_mix, width='stretch')
            with c4:
                res_emp = df_emp_base[(df_emp_base['Empresa'].isin(ordem)) & df_emp_base['Dias p/ Resolver'].notna()]
                fig_box_emp = px.box(res_emp, x='Dias p/ Resolver', y='Empresa', points=False, orientation='h',
                                     category_orders={'Empresa': ordem[::-1]},
                                     title='Dias para resolver, por empresa')
                fig_box_emp.update_layout(height=altura_barras(len(ordem)))
                encurtar_nomes(fig_box_emp, ordem)
                st.plotly_chart(fig_box_emp, width='stretch')

            st.subheader("📉 Evolução mensal por empresa")
            empresas_hist = st.multiselect("Empresas para comparar:", options=ordem, default=ordem[:5],
                                           placeholder="Selecione")
            if empresas_hist:
                base_h = df_emp_base[df_emp_base['Empresa'].isin(empresas_hist)]
                res_h = base_h[base_h['Categoria'] == 'Resolvido']
                if periodo_ini is not None:
                    res_h = res_h[res_h['Mês Fechamento'].between(periodo_ini, periodo_fim)]
                c5, c6 = st.columns(2)
                with c5:
                    serie = res_h.groupby(['Mês Fechamento', 'Empresa']).size().reset_index(name='Resolvidos')
                    fig_r = px.line(serie, x='Mês Fechamento', y='Resolvidos', color='Empresa', markers=True,
                                    title='Chamados resolvidos por mês')
                    fig_r.update_xaxes(tickformat='%m/%Y')
                    st.plotly_chart(fig_r, width='stretch')
                with c6:
                    serie_t = res_h.groupby(['Mês Fechamento', 'Empresa'])['Dias p/ Resolver'].median().reset_index()
                    fig_t = px.line(serie_t, x='Mês Fechamento', y='Dias p/ Resolver', color='Empresa', markers=True,
                                    title='Mediana de dias para resolver')
                    fig_t.update_xaxes(tickformat='%m/%Y')
                    st.plotly_chart(fig_t, width='stretch')

                if 'Tipo' in base_h.columns:
                    calor = resumo_por_grupo(base_h, ['Empresa', 'Tipo'], data_ref)
                    calor = calor.pivot(index='Empresa', columns='Tipo', values='% Resolução')
                    fig_heat = px.imshow(calor, text_auto='.0f', color_continuous_scale='RdYlGn', zmin=0, zmax=100,
                                         aspect='auto', title='% de resolução: empresa × tipo')
                    encurtar_nomes(fig_heat, calor.index)
                    st.plotly_chart(fig_heat, width='stretch')

            st.subheader("📋 Tabela de indicadores por empresa")
            st.dataframe(emp, width='stretch', hide_index=True)
            st.download_button("📥 Baixar indicadores por empresa (CSV)",
                               data=emp.to_csv(index=False, sep=';').encode('utf-8-sig'),
                               file_name='eficiencia_empresas.csv', mime='text/csv')

    # ------------------------------------------------------------ Resumo Executivo
    with aba4:
        st.subheader("🧾 Resumo Executivo")

        hist = historico_mensal(df_final, periodo_ini, periodo_fim)
        emp_all = resumo_por_grupo(df_final, 'Empresa', data_ref)
        emp_all = emp_all[emp_all['Total'] >= 20].copy()
        if not emp_all.empty:
            emp_all['Índice de Eficiência'] = calcular_score(emp_all)
        tipo_res = resumo_por_grupo(df_final, 'Tipo', data_ref) if 'Tipo' in df_final.columns else pd.DataFrame()
        est_res = resumo_por_grupo(df_final, 'Estação', data_ref)

        abertos = df_final[df_final['Categoria'] == 'Aberto'].copy()
        abertos['Idade (dias)'] = (data_ref - abertos['Data']).dt.days

        destaques = comparar_trimestres(hist, data_ref)

        if not abertos.empty:
            mais90 = int((abertos['Idade (dias)'] > 90).sum())
            destaques.append(f"Existem **{len(abertos)}** chamados em aberto; **{mais90}** "
                             f"({pct(mais90, len(abertos)):.0f}%) têm mais de 90 dias.")
        if not emp_all.empty and len(emp_all) > 1:
            e = emp_all.sort_values('Índice de Eficiência', ascending=False)
            destaques.append(f"Empresa mais eficiente (≥20 chamados): **{e.iloc[0]['Empresa']}** "
                             f"(índice {e.iloc[0]['Índice de Eficiência']:.1f}); menos eficiente: "
                             f"**{e.iloc[-1]['Empresa']}** (índice {e.iloc[-1]['Índice de Eficiência']:.1f}).")
            maior_backlog = emp_all.sort_values('Abertos', ascending=False).iloc[0]
            destaques.append(f"Maior backlog: **{maior_backlog['Empresa']}** com {maior_backlog['Abertos']} "
                             f"chamados em aberto ({maior_backlog['Backlog +90 dias']} com +90 dias).")
        if not tipo_res.empty and tipo_res['Mediana Dias Resolução'].notna().any():
            lento = tipo_res.sort_values('Mediana Dias Resolução', ascending=False).iloc[0]
            destaques.append(f"Tipo mais lento para resolver: **{lento['Tipo']}** "
                             f"(mediana {lento['Mediana Dias Resolução']:.0f} dias).")
        reinc = est_res[est_res['Total'] >= 5]
        destaques.append(f"**{len(reinc)}** estações tiveram 5 ou mais chamados no período (reincidência).")
        if 'Dias p/ Acionar' in df_final.columns and df_final['Dias p/ Acionar'].notna().any():
            destaques.append(f"Tempo mediano entre abertura e acionamento da empresa: "
                             f"**{df_final['Dias p/ Acionar'].median():.0f} dia(s)**.")

        st.markdown("\n".join(f"- {d}" for d in destaques))
        st.divider()

        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Resumo por tipo de chamado**")
            if not tipo_res.empty:
                st.dataframe(tipo_res[['Tipo', 'Total', 'Resolvidos', 'Abertos', 'Cancelados', '% Resolução',
                                       '% SLA No Prazo', 'Mediana Dias Resolução']],
                             width='stretch', hide_index=True)
        with c2:
            st.markdown("**Top 15 problemas mais recorrentes**")
            if 'Descrição' in df_final.columns:
                desc = resumo_por_grupo(df_final, 'Descrição', data_ref).sort_values('Total', ascending=False).head(15)
                st.dataframe(desc[['Descrição', 'Total', 'Abertos', '% Resolução', 'Mediana Dias Resolução']],
                             width='stretch', hide_index=True)

        if 'Descrição' in df_final.columns and 'Tipo' in df_final.columns:
            tree = df_final.groupby(['Tipo', 'Descrição']).size().reset_index(name='Chamados')
            fig_tree = px.treemap(tree, path=[px.Constant('Todos'), 'Tipo', 'Descrição'], values='Chamados',
                                  title='Mapa de problemas: tipo → descrição')
            st.plotly_chart(fig_tree, width='stretch')

        c3, c4 = st.columns(2)
        with c3:
            st.markdown("**Estações mais reincidentes**")
            top_est = est_res.sort_values('Total', ascending=False).head(20)
            st.dataframe(top_est[['Estação', 'Total', 'Abertos', 'Resolvidos', '% Resolução',
                                  'Idade Média Backlog (dias)']], width='stretch', hide_index=True)
        with c4:
            st.markdown("**Chamados abertos há mais tempo**")
            cols_crit = [c for c in ['ID Sopro', 'Data', 'Estação', 'Tipo', 'Descrição', 'Empresa', 'Status',
                                     'Idade (dias)'] if c in abertos.columns]
            criticos = abertos.sort_values('Idade (dias)', ascending=False)[cols_crit].head(20)
            st.dataframe(criticos, width='stretch', hide_index=True)

        st.divider()
        resumo_kpis = pd.DataFrame({
            'Indicador': ['Período inicial', 'Período final', 'Total de chamados', 'Em aberto', 'Resolvidos',
                          'Cancelados', '% SLA no prazo', 'Mediana dias p/ resolver', 'Backlog com +90 dias'],
            'Valor': [df_final['Data'].min().strftime('%d/%m/%Y'), df_final['Data'].max().strftime('%d/%m/%Y'),
                      len(df_final), len(abertos), int((df_final['Categoria'] == 'Resolvido').sum()),
                      int((df_final['Categoria'] == 'Cancelado').sum()),
                      f"{pct((df_final['SLA'] == 'NO PRAZO').sum(), df_final['SLA'].isin(['NO PRAZO', 'FORA DO PRAZO']).sum()):.1f}%",
                      df_final['Dias p/ Resolver'].median(), int((abertos['Idade (dias)'] > 90).sum())],
        })
        resumo_kpis['Valor'] = resumo_kpis['Valor'].astype(str)
        destaques_df = pd.DataFrame({'Destaques': [d.replace('**', '') for d in destaques]})
        hist_xlsx = hist.copy()
        if not hist_xlsx.empty:
            hist_xlsx['Mês'] = hist_xlsx['Mês'].dt.strftime('%m/%Y')

        excel = gerar_excel({
            'Resumo': resumo_kpis,
            'Destaques': destaques_df,
            'Histórico Mensal': hist_xlsx,
            'Por Empresa': emp_all.sort_values('Índice de Eficiência', ascending=False) if not emp_all.empty else emp_all,
            'Por Tipo': tipo_res,
            'Por Estação': est_res.sort_values('Total', ascending=False),
            'Backlog em Aberto': abertos.sort_values('Idade (dias)', ascending=False),
            'Dados Filtrados': df_final,
        })
        st.download_button("📥 Baixar Relatório Completo (Excel)", data=excel,
                           file_name=f"relatorio_sopro_{data_ref.strftime('%Y%m%d')}.xlsx",
                           mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # ------------------------------------------------------------- Ninhos / Zeladoria
    colunas_det = ['ID Sopro', 'Data', 'Estação', 'Tipo', 'Descrição', 'Empresa', 'Status', 'Data Fechamento']
    # Estas abas usam período, regional e empresa; os demais filtros não se aplicam (avisa quando estão ativos)
    ignorados = [nome for nome, ativo in [
        ("status", status_selecionado != "TODOS"), ("SLA", sla_selecionado != "TODOS"),
        ("tipo", 'Tipo' in df.columns and set(tipos_selecionados) != set(tipos_disponiveis)),
        ("estações", bool(estacoes_selecionadas))] if ativo]
    aviso_ignorados = (f"Os filtros de {', '.join(ignorados)} da barra lateral não se aplicam a esta aba: ela mostra "
                       "todos os casos do período, da regional e da empresa escolhidos.") if ignorados else None

    with aba5:
        st.subheader("🚨 Sites com Maior Incidência de Ninhos")
        st.write("Sites que mais geram chamado de ninhos")
        if aviso_ignorados:
            st.info(aviso_ignorados)
        ranking_ocorrencias(
            df_f1, lambda d: d['Descrição'].astype(str).str.contains('Insetos Ou Ninhos', case=False, na=False),
            "Ninhos", 'Reds', colunas_det, {"Últimos 3 meses": 3, "Últimos 6 meses": 6, "Últimos 12 meses": 12},
            'periodo_ninhos'
        )

    with aba6:
        st.subheader("🌿 Sites com Maior Incidência de Problemas de Zeladoria")
        st.write("Monitore as estações que mais demandaram serviços de zeladoria no histórico recente.")
        if aviso_ignorados:
            st.info(aviso_ignorados)
        ranking_ocorrencias(
            df_f1, lambda d: d['Tipo'].astype(str).str.contains('Zeladoria', case=False, na=False),
            "Zeladoria", 'Greens', colunas_det, {"Últimos 3 meses": 3, "Últimos 6 meses": 6, "Últimos 12 meses": 12},
            'periodo_zel'
        )
